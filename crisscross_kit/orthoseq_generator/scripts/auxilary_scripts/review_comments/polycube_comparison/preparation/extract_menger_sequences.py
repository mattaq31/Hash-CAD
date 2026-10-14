"""Extract the published Menger handles from Tables S18/S22 without redesigning or filtering their cores."""

from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
from pathlib import Path
import re

from openpyxl.styles import Alignment, Font, PatternFill
import pandas as pd
from pypdf import PdfReader

from orthoseq_generator.sequence_generation import revcom, sorted_key


if __name__ == "__main__":
    folder = Path(__file__).resolve().parents[1]
    source_pdf = folder / "sources" / "diep_polycubes_supplementary.pdf"
    data_dir = folder / "data"
    output_xlsx = data_dir / "menger_published_sequences.xlsx"
    source_url = (
        "https://www.cambridge.org/engage/api-gateway/coe/assets/orp/resource/item/"
        "6942549ce3cb457e131a20ae/original/supplementary-material.pdf"
    )
    doi = "10.26434/chemrxiv-2025-qbk3f"
    spacer = "T" * 12
    table_specs = [
        ("S18", range(91, 93), "S4.3", "Table S18"),
        ("S22", range(100, 110), "S4.7", "Table S22"),
    ]
    row_pattern = re.compile(
        r"^\s*(?P<label>(?:CT(?P<cube>\d+)|no\s*SEs)\s*"
        r"(?P<face>FRONT|BOT|LEFT|TOP|RIGHT|BACK)\s*(?P<strand>[0-3]))\s+(?P<sequence>[ACGT]+)\s*$"
    )
    reader = PdfReader(source_pdf)
    tables = {}
    for table, pages, start_heading, end_caption in table_specs:
        rows = []
        active = False
        finished = False
        for page_number in pages:
            page_text = reader.pages[page_number - 1].extract_text(extraction_mode="layout")
            for line in page_text.splitlines():
                stripped = line.strip()
                if stripped.startswith(start_heading):
                    active = True
                    continue
                if not active:
                    continue
                if re.match(end_caption.replace(" ", r"\s+") + r"\s*:", stripped):
                    finished = True
                    active = False
                    break
                match = row_pattern.fullmatch(line)
                if match:
                    fields = match.groupdict()
                    face = fields["face"]
                    strand_index = int(fields["strand"])
                    cube = int(fields["cube"]) if fields["cube"] is not None else None
                    label_prefix = f"CT{cube}" if cube is not None else "no_SEs"
                    rows.append({
                        "label": f"{label_prefix}_{face}_{strand_index}",
                        "source_label_text": fields["label"],
                        "cube": cube,
                        "face": face,
                        "strand_index": strand_index,
                        "full_sequence": fields["sequence"],
                        "source_table": table,
                        "source_page": page_number,
                    })
                elif re.fullmatch(r"[ACGT]+", stripped):
                    if not rows or rows[-1]["source_page"] != page_number:
                        raise ValueError(f"Unattached sequence continuation in {table}, page {page_number}.")
                    rows[-1]["full_sequence"] += stripped
                elif stripped and stripped != "Label" and not stripped.startswith("Label "):
                    if not stripped.isdigit():
                        raise ValueError(f"Unrecognized table content in {table}, page {page_number}: {line!r}")
            if finished:
                break
        if not finished or not rows:
            raise ValueError(f"Could not locate the complete {table} table in its expected pages.")
        if len({row["label"] for row in rows}) != len(rows):
            raise ValueError(f"Duplicate handle labels in {table}.")
        tables[table] = rows

    binding_rows = tables["S18"]
    bindings = {(row["face"], row["strand_index"]): row for row in binding_rows}
    if len(bindings) != 24:
        raise ValueError(f"Expected 24 scaffold-binding references in S18, extracted {len(bindings)}.")

    handles = tables["S22"]
    for row in handles:
        binding_row = bindings[(row["face"], row["strand_index"])]
        binding = binding_row["full_sequence"]
        if not row["full_sequence"].startswith(binding):
            raise ValueError(f"{row['label']} does not begin with its S18 scaffold-binding sequence.")
        exposed = row["full_sequence"][len(binding):]
        if not exposed.startswith(spacer):
            raise ValueError(f"{row['label']} does not have the expected T12 spacer after its binding segment.")
        # Remove exactly the known spacer length: a sticky-end core can itself start with T.
        core = exposed[len(spacer):]
        if not core:
            raise ValueError(f"{row['label']} has no sticky-end core.")
        row.update({
            "origami_binding_sequence": binding,
            "binding_reference_label": binding_row["label"],
            "binding_reference_table": "S18",
            "binding_reference_page": binding_row["source_page"],
            "spacer": spacer,
            "spacer_length": len(spacer),
            "core": core,
            "core_length": len(core),
            "core_gc_count": core.count("G") + core.count("C"),
            "exposed_sequence": exposed,
            "source_doi": doi,
        })
        if binding + spacer + core != row["full_sequence"]:
            raise ValueError(f"Reconstruction failed for {row['label']}.")

    cores = {row["core"] for row in handles}
    missing_complements = sorted(core for core in cores if revcom(core) not in cores)
    if missing_complements:
        raise ValueError(f"Published cores lack partners: {missing_complements}")
    if any(core == revcom(core) for core in cores):
        raise ValueError("A self-complementary core requires separate pairing semantics.")
    pair_keys = sorted({sorted_key(core, revcom(core)) for core in cores})
    pair_ids = {key: idx for idx, key in enumerate(pair_keys)}
    core_sources = defaultdict(list)
    for row in handles:
        key = sorted_key(row["core"], revcom(row["core"]))
        row["global_pair_id"] = pair_ids[key]
        row["pair_slot"] = "H" if row["core"] == key[0] else "A"
        core_sources[row["core"]].append(row)

    pairs = []
    for idx, (core, rc_core) in enumerate(pair_keys):
        pairs.append({
            "pair_idx": idx,
            "global_pair_id": idx,
            "seq": spacer + core,
            "rc_seq": spacer + rc_core,
            "core": core,
            "rc_core": rc_core,
            "spacer": spacer,
            "seq_source_labels": ";".join(row["label"] for row in core_sources[core]),
            "rc_seq_source_labels": ";".join(row["label"] for row in core_sources[rc_core]),
            "seq_source_pages": ";".join(map(str, sorted({row["source_page"] for row in core_sources[core]}))),
            "rc_seq_source_pages": ";".join(map(str, sorted({row["source_page"] for row in core_sources[rc_core]}))),
            "seq_occurrences": len(core_sources[core]),
            "rc_seq_occurrences": len(core_sources[rc_core]),
        })

    # These checks audit the extraction; no entries are dropped to make the counts agree.
    counts = (len(handles), len(cores), len(pairs))
    if counts != (216, 192, 96):
        raise ValueError(f"Unexpected handle/core/pair counts: {counts}; inspect the source before proceeding.")
    if {row["core_length"] for row in handles} != {9}:
        raise ValueError("Unexpected core lengths; inspect the published sequences before proceeding.")
    face_counts = Counter((row["cube"], row["face"]) for row in handles)
    if set(face_counts.values()) != {4} or {row["cube"] for row in handles} != set(range(19)):
        raise ValueError("Incomplete four-strand faces or cube coverage in S22.")

    metadata = {
        "source.doi": doi,
        "source.supplementary_url": source_url,
        "source.pdf": str(source_pdf.relative_to(folder)),
        "source.pdf_sha256": hashlib.sha256(source_pdf.read_bytes()).hexdigest(),
        "source.table": "S22",
        "source.pages": "100-109 (1-based PDF pages; same as printed pages)",
        "source.binding_reference": "S18, pages 91-92",
        "source.reported_core_length_S1_8": 8,
        "source.reported_tm_celsius_S1_8": 15,
        "source.reported_tm_tolerance_celsius_S1_8": 0.5,
        "source.reported_gc_fraction_S1_8": "5/8",
        "extraction.created_at_utc": datetime.now(timezone.utc).isoformat(timespec="seconds"),
        "extraction.method": "pypdf layout rows; join DNA-only continuations; match S18 prefix; remove exactly T12",
        "extraction.published_handle_count": len(handles),
        "extraction.unique_core_count": len(cores),
        "found_pair_count": len(pairs),
        "input.source_kind": "published_polycube_menger_library",
        "input.length": 9,
        "input.fivep_ext": spacer,
        "input.threep_ext": "",
        "verified_with_direct_nupack": False,
        "extraction.length_discrepancy": "S1.8 reports L=8; every extracted S22 binding core has length 9.",
        "extraction.orientation": "All sequences are 5'-3'; both pair slots carry their own 5' T12 spacer.",
        "extraction.deduplication": "Unique core/complement pairs; original occurrences retained in source_handles.",
        "extraction.label_normalization": "CT{cube}_{face}_{strand}; source_label_text retains extracted label text.",
    }
    data_dir.mkdir(parents=True, exist_ok=True)
    handles_df = pd.DataFrame(handles)
    with pd.ExcelWriter(output_xlsx, engine="openpyxl") as writer:
        pd.DataFrame(pairs).to_excel(writer, sheet_name="found_pairs", index=False)
        handles_df.to_excel(writer, sheet_name="source_handles", index=False)
        pd.DataFrame(binding_rows).to_excel(writer, sheet_name="binding_references", index=False)
        pd.DataFrame(metadata.items(), columns=["key", "value"]).to_excel(
            writer, sheet_name="run_metadata", index=False,
        )
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "C2" if sheet.title != "run_metadata" else "A2"
            sheet.auto_filter.ref = sheet.dimensions
            sheet.sheet_view.showGridLines = False
            for column in sheet.columns:
                width = min(85, max(18, max(len(str(cell.value or "")) for cell in column) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
                for cell in column:
                    cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
                    cell.alignment = Alignment(vertical="top", wrap_text=len(str(cell.value or "")) > width)
                column[0].fill = PatternFill("solid", fgColor="243746")
                column[0].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            sheet.row_dimensions[1].height = 30
            for row in sheet.iter_rows(min_row=2):
                if any(cell.alignment.wrap_text for cell in row):
                    sheet.row_dimensions[row[0].row].height = 42

    print(f"Extracted {len(handles)} handles, {len(cores)} unique 9-nt cores, and {len(pairs)} complementary pairs.")
    print(f"Verified T12 spacers, S18 prefixes, sequence reconstruction, and all {len(face_counts)} four-strand faces.")
    print(f"Saved {output_xlsx}")
    print("Discrepancy retained: S1.8 reports L=8, but the published S22 cores are all 9 nt.")
