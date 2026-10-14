"""Evaluate both complete libraries at one model and save association and complex free energy comparisons."""

import hashlib
from pathlib import Path

import numpy as np
import pandas as pd
from nupack import Model, __version__ as nupack_version
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from crisscross.helper_functions.lab_helper_sheet_generation import adjust_column_width
from orthoseq_generator.search_report_reader import load_found_pairs, load_metadata
from orthoseq_generator.search_reporting import (
    build_selected_sequence_data, verify_selected_pairs, write_hybrid_search_result_xlsx,
)
from sequence_quality import association_to_complex_energies, compute_specificity_gaps


if __name__ == "__main__":
    folder = Path(__file__).resolve().parent
    nupack_params = {
        "material": "dna", "celsius": 37.0, "sodium": 0.05, "magnesium": 0.0125, "energy_type": "total",
    }
    data_dir = folder / "data"
    inputs = {
        "Menger": data_dir / "menger_published_sequences.xlsx",
        "OrthoSeq": data_dir / "orthoseq_assembly_search_2026-10-08_19-51-55.xlsx",
    }
    source_commit = "9e0203cda2afac87399687f906f35eec8f6e0f01"
    model = Model(**{key: value for key, value in nupack_params.items() if key != "energy_type"})
    comparison_rows = []
    specificity_rows = []
    libraries = {}
    for library, input_xlsx in inputs.items():
        pairs = load_found_pairs(input_xlsx)
        source_metadata = load_metadata(input_xlsx)
        count = len(pairs)
        sequence_pairs = list(zip(pairs["seq"], pairs["rc_seq"]))
        sequences = [sequence for pair in sequence_pairs for sequence in pair]
        if count < 2 or len(set(sequences)) != 2 * count:
            raise ValueError(f"{library} must contain at least two distinct exposed sequence pairs.")
        if any(len(sequence) != 21 or not sequence.startswith("T" * 12) for sequence in sequences):
            raise ValueError(f"{library} must contain 9-nt cores with their own 5' T12 spacers.")
        # Retain published handle provenance while using the standard builder for our search export.
        selected = pairs.to_dict("records") if library == "Menger" else build_selected_sequence_data(
            sequence_pairs, pairs["global_pair_id"].astype(int).tolist(),
        )
        print(f"Evaluating all {count} {library} pairs without filtering: {model}", flush=True)
        # Each verification resets the backend globals and passes the settings to its worker processes.
        energies = verify_selected_pairs(selected, nupack_params=nupack_params)
        try:
            raw = association_to_complex_energies(energies)
            raw_basis = "Recovered from uncapped association and monomer energies."
        except RuntimeError:
            print(f"Computing {library} complex free energies with the shared backend.", flush=True)
            raw = verify_selected_pairs(selected, nupack_params={**nupack_params, "energy_type": "totalu"})
            raw_basis = "Direct shared-backend totalu calculation, including the usual -1 kcal/mol cap."
        scores = {
            "total": compute_specificity_gaps(energies["on_target_energies"], energies["off_target"]),
            "totalu": compute_specificity_gaps(raw["on_target_energies"], raw["off_target"]),
        }
        on = np.asarray(energies["on_target_energies"], dtype=float)
        self_energies = np.concatenate([energies["self_energy_seqs"], energies["self_energy_rc_seqs"]])
        off = energies["off_target"]
        lower = np.tril_indices(count, -1)
        cross_mask = ~np.eye(count, dtype=bool)
        cross = np.concatenate([
            off["handle_handle_energies"][lower], off["antihandle_handle_energies"][cross_mask],
            off["antihandle_antihandle_energies"][lower],
        ])
        homo = np.concatenate([np.diag(off["handle_handle_energies"]), np.diag(off["antihandle_antihandle_energies"])])
        if cross.size != 2 * count * (count - 1):
            raise ValueError(f"Unexpected number of distinct cross-pair interactions for {library}.")
        if not all(np.all(np.isfinite(values)) for values in (on, self_energies, cross, homo)):
            raise ValueError(f"Non-finite energies for {library}.")
        comparison = {
            "library": library, "pair_count": count,
            "on_target_min_kcal_mol": float(on.min()), "on_target_max_kcal_mol": float(on.max()),
            "on_target_mean_kcal_mol": float(on.mean()), "cross_pair_count": int(cross.size),
            "cross_pair_min_kcal_mol": float(cross.min()), "homodimer_min_kcal_mol": float(homo.min()),
            "self_folding_min_kcal_mol": float(self_energies.min()),
            "specificity_gap_total_min_kcal_mol": scores["total"]["minimum_gap"],
            "specificity_gap_total_mean_kcal_mol": scores["total"]["mean_gap"],
            "specificity_gap_totalu_min_kcal_mol": scores["totalu"]["minimum_gap"],
            "specificity_gap_totalu_mean_kcal_mol": scores["totalu"]["mean_gap"],
        }
        comparison_rows.append(comparison)
        strand_scores = []
        for convention, quality in scores.items():
            for strand_idx, gap in enumerate(quality["strand_gaps"]):
                competitor_idx = int(quality["strongest_offtarget_indices"][strand_idx])
                row, competitor = pairs.iloc[strand_idx % count], pairs.iloc[competitor_idx % count]
                slot, off_slot = ("H" if strand_idx < count else "A"), ("H" if competitor_idx < count else "A")
                column, off_column = ("seq" if slot == "H" else "rc_seq"), ("seq" if off_slot == "H" else "rc_seq")
                strand_scores.append({
                    "library": library, "energy_type": convention,
                    "global_pair_id": int(row["global_pair_id"]), "strand_slot": slot, "sequence": row[column],
                    "source_labels": row.get(f"{column}_source_labels", ""),
                    "specificity_gap_kcal_mol": float(gap),
                    "on_target_energy_kcal_mol": float(quality["on_target_energies"][strand_idx]),
                    "strongest_offtarget_energy_kcal_mol": float(quality["strongest_offtarget_energies"][strand_idx]),
                    "offtarget_pair_id": int(competitor["global_pair_id"]), "offtarget_slot": off_slot,
                    "offtarget_sequence": competitor[off_column],
                    "offtarget_source_labels": competitor.get(f"{off_column}_source_labels", ""),
                    "offtarget_is_homodimer": strand_idx == competitor_idx,
                })
        specificity_rows.extend(strand_scores)
        homodimers = [
            {"global_pair_id": int(row["global_pair_id"]), "strand_slot": slot,
             "sequence": row[column], "association_energy_kcal_mol": float(off[matrix][idx, idx])}
            for idx, row in pairs.iterrows()
            for slot, column, matrix in [
                ("H", "seq", "handle_handle_energies"), ("A", "rc_seq", "antihandle_antihandle_energies"),
            ]
        ]
        strongest = None
        for matrix_name, row_slot, column_slot in [
            ("handle_handle_energies", "H", "H"), ("antihandle_handle_energies", "H", "A"),
            ("antihandle_antihandle_energies", "A", "A"),
        ]:
            positions = np.argwhere(cross_mask if row_slot != column_slot else np.tril(cross_mask))
            values = off[matrix_name][positions[:, 0], positions[:, 1]]
            row_idx, column_idx = positions[int(np.argmin(values))]
            energy = float(off[matrix_name][row_idx, column_idx])
            if strongest is None or energy < strongest["energy_kcal_mol"]:
                strongest = {"energy_kcal_mol": energy, "interaction": matrix_name}
                for participant, idx, slot in [("a", row_idx, row_slot), ("b", column_idx, column_slot)]:
                    row = pairs.iloc[idx]
                    column = "seq" if slot == "H" else "rc_seq"
                    strongest.update({
                        f"pair_{participant}": int(row["global_pair_id"]), f"slot_{participant}": slot,
                        f"sequence_{participant}": row[column],
                        f"source_labels_{participant}": row.get(f"{column}_source_labels", ""),
                    })
        summary = [{"metric": key, "value": value} for key, value in comparison.items() if key != "library"]
        summary.extend([
            {"metric": "on_target_std_population_kcal_mol", "value": float(on.std())},
            {"metric": "cross_pair_mean_kcal_mol", "value": float(cross.mean())},
            {"metric": "homodimer_count", "value": int(homo.size)},
            {"metric": "self_folding_mean_kcal_mol", "value": float(self_energies.mean())},
        ])
        metadata = {key: value for key, value in source_metadata.items()
                    if key.startswith(("source.", "extraction."))}
        metadata.update({
            "source.input_xlsx": str(input_xlsx),
            "source.input_sha256": hashlib.sha256(input_xlsx.read_bytes()).hexdigest(),
            "evaluation.settings_basis": "Common assembly-buffer model; original authors' design run unconfirmed.",
            "evaluation.temperature_basis": "Chosen 37 C comparison temperature within the assembly annealing window.",
            "evaluation.monovalent_salt_basis": "50 mM approximation; 1x TAE does not specify NUPACK sodium.",
            "evaluation.magnesium_basis": "Reported nominal 12.5 mM; not a reconstruction of free-ion concentration.",
            "evaluation.experimental_buffer_reported": "1x TAE + 12.5 mM Mg2+; SI S2.1, page 35.",
            "evaluation.experimental_menger_annealing": "Slow ramp 44 to 28.5 C; cool to and hold at 20 C; SI page 36.",
            "evaluation.experimental_buffer_exactly_reproduced": False,
            "evaluation.menger_original_matrix_settings_confirmed": False,
            "evaluation.source_code_commit": source_commit,
            "evaluation.source_code_url": f"https://github.com/sulcgroup/DNA3Dpolycubes/tree/{source_commit}/SEdesign",
            "evaluation.filtered_or_redesigned": False,
            "evaluation.specificity_competitors": "All library strands including homodimers; exclude intended partner.",
            "evaluation.raw_complex_basis": raw_basis,
            "evaluation.energy_convention": "Association DeltaG = G_AB - G_A - G_B; complex free energy G_AB.",
            "evaluation.weak_energy_clipping": "Shared backend caps total and totalu values above -1 kcal/mol at -1.",
            "nupack.version": nupack_version, "nupack.ensemble": "stacking", "nupack.resolved_model": str(model),
        })
        metadata.update({f"source.original_{key}": value for key, value in source_metadata.items()
                         if key.startswith(("nupack.", "search."))})
        extra_sheets = {"homodimers": homodimers, "energy_summary": summary, "strongest_cross_pair": [strongest]}
        if library == "Menger":
            extra_sheets["source_handles"] = pd.read_excel(input_xlsx, sheet_name="source_handles").to_dict("records")
        libraries[library] = {
            "selected": selected, "verified": energies, "metadata": metadata, "source_metadata": source_metadata,
            "input_xlsx": input_xlsx, "count": count, "scores": strand_scores, "extra_sheets": extra_sheets,
        }

    # Write the standard report once per library; formatting and comparison logic stay local to this script.
    data_dir.mkdir(parents=True, exist_ok=True)
    for library, data in libraries.items():
        output_name = f"orthoseq_{data['count']}_energy_evaluation.xlsx"
        if library == "Menger":
            output_name = "menger_energy_evaluation.xlsx"
        output_xlsx = data_dir / output_name
        extra_sheets = {**data["extra_sheets"], "comparison": comparison_rows,
                        "specificity_gaps": data["scores"] if library == "Menger" else specificity_rows}
        metadata = {**data["metadata"], "evaluation.library_sizes_matched":
                    libraries["Menger"]["count"] == libraries["OrthoSeq"]["count"]}
        write_hybrid_search_result_xlsx(
            output_xlsx, algorithm_name="complete_library_common_model_evaluation",
            selected_sequence_data=data["selected"], verified=data["verified"], nupack_params=nupack_params,
            search_params={}, input_params={key.removeprefix("input."): value for key, value in
                                           data["source_metadata"].items() if key.startswith("input.")},
            artifact_info={"dataset_dir": str(data["input_xlsx"].parent)}, generation_data=[],
            validation_data=[{"check": "unfiltered_input_pairs_evaluated", "count": data["count"]}],
            extra_metadata=metadata, extra_sheets=extra_sheets,
        )
        workbook = load_workbook(output_xlsx)
        for sheet in workbook.worksheets:
            adjust_column_width(sheet)
            sheet.freeze_panes = "B2" if sheet.title.startswith("selected_") else "A2"
            sheet.sheet_view.showGridLines = False
            if sheet.max_row > 1:
                sheet.auto_filter.ref = sheet.dimensions
            for column in sheet.columns:
                width = min(85, max(18, sheet.column_dimensions[column[0].column_letter].width + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
                for cell in column:
                    cell.font = Font(name="Arial", size=10, bold=cell.row == 1)
                    cell.alignment = Alignment(vertical="top", wrap_text=True)
                    if cell.row > 1 and isinstance(cell.value, float):
                        cell.number_format = "0" if cell.value.is_integer() else "0.000000"
                column[0].fill = PatternFill("solid", fgColor="243746")
                column[0].font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
            sheet.row_dimensions[1].height = 30
            for row in sheet.iter_rows(min_row=2):
                lines = max(len(str(cell.value or "")) // int(sheet.column_dimensions[cell.column_letter].width) + 1
                            for cell in row)
                sheet.row_dimensions[row[0].row].height = max(18, 16 * lines)
        workbook.save(output_xlsx)
        print(f"Saved {output_xlsx}", flush=True)
    print(pd.DataFrame(comparison_rows).to_string(index=False), flush=True)
